from pathlib import Path
import re
import shutil

from openpyxl import load_workbook

from app.schemas import FNADataPayload
from datetime import datetime


# ---------------------------------------------------------
# Paths
# ---------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parents[1]

TEMPLATE_PATH = (
    BASE_DIR
    / "templates"
    / "2025-01 02 FNA INFO COLLECT TEMPLATE.xlsx"
)

OUTPUT_DIR = BASE_DIR / "data" / "processed"


# ---------------------------------------------------------
# Filename helper
# ---------------------------------------------------------

def safe_filename(name: str | None) -> str:
    """
    Convert a client name into a safe Windows filename.
    """

    name = name or "Client"

    cleaned = re.sub(
        r"[^A-Za-z0-9 _-]",
        "",
        name,
    ).strip()

    return cleaned.replace(" ", "_") or "Client"


# ---------------------------------------------------------
# Name helper
# ---------------------------------------------------------

def split_full_name(full_name: str | None):
    """
    Split the API full_name field into the separate
    name fields used by the Royal Square FNA workbook.

    Adviser verification is still required.
    """

    if not full_name:
        return None, None, None

    parts = full_name.strip().split()

    if len(parts) == 1:
        return parts[0], None, None

    if len(parts) == 2:
        return parts[0], None, parts[1]

    first_name = parts[0]
    second_name = " ".join(parts[1:-1])
    surname = parts[-1]

    return first_name, second_name, surname


# ---------------------------------------------------------
# Calculations
# ---------------------------------------------------------

def calculate_fna_summary(
    payload: FNADataPayload,
) -> dict[str, float]:

    all_assets = (
        payload.assets.properties
        + payload.assets.vehicles
        + payload.assets.savings
        + payload.assets.unit_trusts
    )

    all_liabilities = (
        payload.liabilities.mortgages
        + payload.liabilities.vehicle_finance
        + payload.liabilities.personal_loans
        + payload.liabilities.credit_cards
    )

    all_expenses = (
        payload.household_expenses.fixed_expenses
        + payload.household_expenses.variable_expenses
    )

    total_assets = sum(
        float(asset.current_value)
        for asset in all_assets
    )

    total_liabilities = sum(
        float(liability.outstanding_balance)
        for liability in all_liabilities
    )

    total_expenses = sum(
        float(expense.monthly_amount)
        for expense in all_expenses
    )

    net_monthly_income = (
        float(
            payload.client_demographics.net_monthly_income
        )
        if payload.client_demographics.net_monthly_income
        is not None
        else 0.0
    )

    return {
        "total_assets": total_assets,
        "total_liabilities": total_liabilities,
        "net_worth": total_assets - total_liabilities,
        "total_household_expenses": total_expenses,
        "monthly_disposable_income": (
            net_monthly_income - total_expenses
        ),
    }


# ---------------------------------------------------------
# Asset writer
# ---------------------------------------------------------

def write_assets(
    sheet,
    items,
    start_row: int,
    end_row: int,
):
    """
    Royal Square template:
    column J = asset description
    column L = current value
    """

    for row_number, item in zip(
        range(start_row, end_row + 1),
        items,
    ):
        sheet[f"J{row_number}"] = item.description

        sheet[f"L{row_number}"] = float(
            item.current_value
        )

        sheet[f"L{row_number}"].number_format = (
            'R #,##0.00'
        )


# ---------------------------------------------------------
# Liability writer
# ---------------------------------------------------------

def write_liabilities(
    sheet,
    items,
    start_row: int,
    end_row: int,
):
    """
    Royal Square template:
    column N = outstanding balance
    column Q = financial institution / creditor
    """

    for row_number, item in zip(
        range(start_row, end_row + 1),
        items,
    ):
        sheet[f"N{row_number}"] = float(
            item.outstanding_balance
        )

        sheet[f"N{row_number}"].number_format = (
            'R #,##0.00'
        )

        if item.institution:
            sheet[f"Q{row_number}"] = (
                item.institution
            )
        elif item.description:
            sheet[f"Q{row_number}"] = (
                item.description
            )


# ---------------------------------------------------------
# Expense mapping
# ---------------------------------------------------------

EXPENSE_CELL_MAP = {
    "rent": "V34",
    "rental": "V34",

    "water": "V35",
    "electricity": "V35",

    "rates": "V36",
    "taxes": "V36",

    "petrol": "V39",
    "transport": "V39",

    "groceries": "V40",
    "grocery": "V40",

    "telephone": "V41",
    "data": "V41",

    "cell phone": "V42",
    "mobile": "V42",

    "gym": "V43",

    "housekeeper": "V44",
    "domestic": "V44",

    "school": "V45",
    "education": "V45",

    "dstv": "V46",

    "security": "V47",

    "donation": "V48",
    "tithe": "V48",

    "hobbies": "V49",
    "club": "V49",

    "shopping": "V50",

    "entertainment": "V51",

    "car insurance": "V52",
    "vehicle insurance": "V52",

    "home insurance": "V53",
}


def write_expenses(
    sheet,
    payload: FNADataPayload,
):
    expenses = (
        payload.household_expenses.fixed_expenses
        + payload.household_expenses.variable_expenses
    )

    for expense in expenses:

        description = (
            expense.description
            .strip()
            .lower()
        )

        target_cell = None

        for keyword, cell in EXPENSE_CELL_MAP.items():
            if keyword in description:
                target_cell = cell
                break

        # Some extracted expenses may not correspond
        # to a predefined Royal Square expense field.
        if target_cell is None:
            continue

        existing_value = sheet[target_cell].value

        if not isinstance(
            existing_value,
            (int, float),
        ):
            existing_value = 0

        sheet[target_cell] = (
            float(existing_value)
            + float(expense.monthly_amount)
        )

        sheet[target_cell].number_format = (
            'R #,##0.00'
        )


# ---------------------------------------------------------
# Main population function
# ---------------------------------------------------------

def populate_fna_workbook(
    payload: FNADataPayload,
):
    """
    Create a populated copy of the Royal Square
    Financial Needs Analysis workbook.

    The original blank template is never modified.
    """

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"FNA template not found: {TEMPLATE_PATH}"
        )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    demographics = payload.client_demographics

    client_name = (
        demographics.full_name
        or "Client"
    )

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")

    output_filename = (
        f"{safe_filename(client_name)}"
        f"_FNA_Processed_{timestamp}.xlsx"
    )

    output_path = (
        OUTPUT_DIR / output_filename
    )

    # Copy the blank workbook first.
    shutil.copy2(
        TEMPLATE_PATH,
        output_path,
    )

    workbook = load_workbook(
        output_path,
        data_only=False,
    )

    if "Info Collect" not in workbook.sheetnames:
        raise ValueError(
            "The workbook does not contain "
            "the 'Info Collect' worksheet."
        )

    sheet = workbook["Info Collect"]

    # -----------------------------------------------------
    # Client information
    # -----------------------------------------------------

    (
        first_name,
        second_name,
        surname,
    ) = split_full_name(
        demographics.full_name
    )

    if first_name:
        sheet["B8"] = first_name

    if second_name:
        sheet["B9"] = second_name

    if surname:
        sheet["B10"] = surname

    if demographics.marital_status:
        sheet["D7"] = (
            demographics.marital_status
        )

    if demographics.id_number:
        sheet["B13"] = (
            demographics.id_number
        )

    if demographics.employer:
        sheet["B18"] = (
            demographics.employer
        )

    if demographics.tax_number:
        sheet["B21"] = (
            demographics.tax_number
        )

    # -----------------------------------------------------
    # Income
    # -----------------------------------------------------

    if (
        demographics.gross_monthly_income
        is not None
    ):
        sheet["T8"] = float(
            demographics.gross_monthly_income
        )

        sheet["T8"].number_format = (
            'R #,##0.00'
        )

    if (
        demographics.net_monthly_income
        is not None
    ):
        sheet["T12"] = float(
            demographics.net_monthly_income
        )

        sheet["T12"].number_format = (
            'R #,##0.00'
        )

    # -----------------------------------------------------
    # Assets
    # -----------------------------------------------------

    write_assets(
        sheet,
        payload.assets.properties,
        7,
        10,
    )

    write_assets(
        sheet,
        payload.assets.vehicles,
        13,
        15,
    )

    write_assets(
        sheet,
        payload.assets.savings,
        20,
        24,
    )

    write_assets(
        sheet,
        payload.assets.unit_trusts,
        27,
        35,
    )

    # -----------------------------------------------------
    # Liabilities
    # -----------------------------------------------------

    write_liabilities(
        sheet,
        payload.liabilities.mortgages,
        7,
        10,
    )

    write_liabilities(
        sheet,
        payload.liabilities.vehicle_finance,
        13,
        15,
    )

    write_liabilities(
        sheet,
        payload.liabilities.credit_cards,
        19,
        21,
    )

    write_liabilities(
        sheet,
        payload.liabilities.personal_loans,
        27,
        33,
    )

    # -----------------------------------------------------
    # Expenses
    # -----------------------------------------------------

    write_expenses(
        sheet,
        payload,
    )

    # -----------------------------------------------------
    # Backend-calculated FNA summary
    # -----------------------------------------------------

    summary = calculate_fna_summary(
        payload
    )

    summary_sheet_name = (
        "Automation Summary"
    )

    if (
        summary_sheet_name
        in workbook.sheetnames
    ):
        del workbook[summary_sheet_name]

    summary_sheet = workbook.create_sheet(
        summary_sheet_name
    )

    summary_sheet["A1"] = (
        "Royal Square FNA Automation Summary"
    )

    summary_sheet["A3"] = "Client"
    summary_sheet["B3"] = client_name

    summary_sheet["A5"] = "Total Assets"
    summary_sheet["B5"] = (
        summary["total_assets"]
    )

    summary_sheet["A6"] = (
        "Total Liabilities"
    )
    summary_sheet["B6"] = (
        summary["total_liabilities"]
    )

    summary_sheet["A7"] = "Net Worth"
    summary_sheet["B7"] = (
        summary["net_worth"]
    )

    summary_sheet["A8"] = (
        "Total Household Expenses"
    )
    summary_sheet["B8"] = (
        summary[
            "total_household_expenses"
        ]
    )

    summary_sheet["A9"] = (
        "Monthly Disposable Income"
    )
    summary_sheet["B9"] = (
        summary[
            "monthly_disposable_income"
        ]
    )

    for cell_reference in [
        "B5",
        "B6",
        "B7",
        "B8",
        "B9",
    ]:
        summary_sheet[
            cell_reference
        ].number_format = 'R #,##0.00'

    workbook.save(
        output_path
    )

    return output_path, summary